import csv
import os
import sqlite3
import tempfile
from pathlib import Path

import pandas as pd
import streamlit as st
from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell

DEFAULT_CHUNK_SIZE = 50_000
XLSX_MAX_DATA_ROWS = 1_048_575  # Excel worksheet row limit minus header row


def _clean_headers(headers):
    return [str(value).strip() if value is not None else "" for value in headers]


def get_columns(uploaded_file):
    """Read headers from CSV or the first worksheet of an XLSX file."""
    suffix = Path(uploaded_file.name).suffix.lower()
    uploaded_file.seek(0)
    if suffix == ".csv":
        frame = pd.read_csv(uploaded_file, nrows=0, dtype=str, keep_default_na=False, encoding_errors="replace")
        return _clean_headers(frame.columns.tolist())
    if suffix == ".xlsx":
        workbook = load_workbook(uploaded_file, read_only=True, data_only=True)
        try:
            sheet = workbook[workbook.sheetnames[0]]
            first_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), None)
            return _clean_headers(list(first_row or []))
        finally:
            workbook.close()
    raise ValueError("Unsupported file type. Please upload CSV or XLSX.")


def iter_csv_chunks(uploaded_file, chunk_size):
    uploaded_file.seek(0)
    yield from pd.read_csv(
        uploaded_file,
        dtype=str,
        keep_default_na=False,
        chunksize=chunk_size,
        encoding_errors="replace",
    )


def iter_xlsx_chunks(uploaded_file, chunk_size):
    """Stream the first worksheet in bounded batches using openpyxl read-only mode."""
    uploaded_file.seek(0)
    workbook = load_workbook(uploaded_file, read_only=True, data_only=True)
    try:
        sheet = workbook[workbook.sheetnames[0]]
        rows = sheet.iter_rows(values_only=True)
        header = next(rows, None)
        if header is None:
            raise ValueError("The Excel workbook is empty.")
        columns = _clean_headers(list(header))
        batch = []
        for row in rows:
            values = list(row)
            if len(values) < len(columns):
                values.extend([None] * (len(columns) - len(values)))
            elif len(values) > len(columns):
                values = values[:len(columns)]
            batch.append(["" if value is None else str(value) for value in values])
            if len(batch) >= chunk_size:
                yield pd.DataFrame(batch, columns=columns)
                batch = []
        if batch:
            yield pd.DataFrame(batch, columns=columns)
    finally:
        workbook.close()


def create_outputs(uploaded_file, selected_column, output_dir, original_name, chunk_size):
    """Deduplicate across the whole input and create CSV + XLSX outputs incrementally."""
    suffix = Path(original_name).suffix.lower()
    stem = Path(original_name).stem
    csv_path = os.path.join(output_dir, f"deduped_{stem}.csv")
    xlsx_path = os.path.join(output_dir, f"deduped_{stem}.xlsx")
    db_path = os.path.join(output_dir, "seen_values.sqlite")

    conn = sqlite3.connect(db_path)
    conn.execute("PRAGMA journal_mode=WAL")
    conn.execute("PRAGMA synchronous=NORMAL")
    conn.execute("CREATE TABLE seen (value TEXT PRIMARY KEY)")
    conn.execute("CREATE TEMP TABLE candidates (value TEXT PRIMARY KEY, row_num INTEGER)")

    xlsx_workbook = Workbook(write_only=True)
    xlsx_sheet_number = 0
    xlsx_sheet = None
    xlsx_rows_in_sheet = 0

    rows_processed = rows_kept = duplicates_removed = chunk_count = 0
    csv_header_written = False
    columns = None

    def new_xlsx_sheet(headers):
        nonlocal xlsx_sheet_number, xlsx_sheet, xlsx_rows_in_sheet
        xlsx_sheet_number += 1
        xlsx_sheet = xlsx_workbook.create_sheet(title=f"Deduped_{xlsx_sheet_number}")
        xlsx_sheet.append(headers)
        xlsx_rows_in_sheet = 0

    try:
        chunk_iter = iter_csv_chunks(uploaded_file, chunk_size) if suffix == ".csv" else iter_xlsx_chunks(uploaded_file, chunk_size)
        with open(csv_path, "w", newline="", encoding="utf-8-sig") as csv_handle:
            csv_writer = None
            for chunk_count, chunk in enumerate(chunk_iter, start=1):
                chunk.columns = _clean_headers(chunk.columns.tolist())
                columns = chunk.columns.tolist()
                if selected_column not in columns:
                    raise ValueError(f"Column '{selected_column}' was not found in the input data.")

                if csv_writer is None:
                    csv_writer = csv.writer(csv_handle)
                    csv_writer.writerow(columns)
                    new_xlsx_sheet(columns)

                original_count = len(chunk)
                # Remove duplicates within the current chunk, keeping the first row.
                chunk = chunk.drop_duplicates(subset=[selected_column], keep="first").copy()
                values = chunk[selected_column].astype(str).tolist()

                conn.execute("DELETE FROM candidates")
                conn.executemany(
                    "INSERT OR IGNORE INTO candidates(value, row_num) VALUES (?, ?)",
                    [(value, index) for index, value in enumerate(values)],
                )
                new_indexes = [
                    row[0] for row in conn.execute(
                        "SELECT c.row_num FROM candidates c LEFT JOIN seen s ON s.value=c.value "
                        "WHERE s.value IS NULL ORDER BY c.row_num"
                    )
                ]
                conn.execute("INSERT OR IGNORE INTO seen(value) SELECT value FROM candidates")
                conn.commit()

                kept = chunk.iloc[new_indexes]
                duplicates_removed += original_count - len(kept)
                rows_processed += original_count
                rows_kept += len(kept)

                if not kept.empty:
                    # CSV output is streamed; no complete output file is held in memory.
                    csv_writer.writerows(kept.itertuples(index=False, name=None))
                    # Excel output rolls over to additional sheets when Excel's row limit is reached.
                    for row in kept.itertuples(index=False, name=None):
                        if xlsx_rows_in_sheet >= XLSX_MAX_DATA_ROWS:
                            new_xlsx_sheet(columns)
                        xlsx_sheet.append([None if value == "" else value for value in row])
                        xlsx_rows_in_sheet += 1

                yield {
                    "chunk": chunk_count,
                    "processed": rows_processed,
                    "kept": rows_kept,
                    "duplicates": duplicates_removed,
                    "csv_path": csv_path,
                    "xlsx_path": xlsx_path,
                }

            if chunk_count == 0:
                raise ValueError("The uploaded file has no data rows or could not be read.")
            xlsx_workbook.save(xlsx_path)
    except Exception:
        for path in (csv_path, xlsx_path):
            if os.path.exists(path):
                os.remove(path)
        raise
    finally:
        conn.close()
        try:
            xlsx_workbook.close()
        except Exception:
            pass

    return {
        "csv_path": csv_path,
        "xlsx_path": xlsx_path,
        "processed": rows_processed,
        "kept": rows_kept,
        "duplicates": duplicates_removed,
        "chunks": chunk_count,
    }


st.set_page_config(page_title="CSV & Excel Duplicate Remover", page_icon="🧹", layout="wide")
st.title("🧹 CSV & Excel Duplicate Remover")
st.write(
    "Upload a CSV or Excel workbook (.xlsx), choose a column to check for duplicates, "
    "and download the cleaned results as CSV, XLSX, or both. The first occurrence is kept."
)
st.info(
    "Large-file mode processes rows in chunks and tracks seen values in a temporary, disk-backed SQLite database. "
    "For XLSX inputs, the first worksheet is processed. XLSX output automatically uses additional worksheets "
    "if the cleaned result exceeds Excel's per-sheet row limit."
)
st.warning(
    "On Streamlit Cloud, uploaded files are transmitted to Streamlit's servers. This app does not upload files to "
    "GitHub or a permanent database, but temporary processing files are created on the server."
)

uploaded_file = st.file_uploader("Select CSV or Excel file", type=["csv", "xlsx"])

if uploaded_file is not None:
    st.caption(f"File: {uploaded_file.name} · Size: {uploaded_file.size / (1024 * 1024):,.2f} MB")
    try:
        columns = get_columns(uploaded_file)
        if not columns:
            st.error("No columns were found. Ensure the first row contains column headers.")
        else:
            selected_column = st.selectbox(
                "Select column to remove duplicates by",
                options=columns,
                help="Rows with repeated values in this column are removed; the first occurrence is kept.",
            )
            chunk_size = st.select_slider(
                "Rows processed per chunk",
                options=[10_000, 25_000, 50_000, 100_000],
                value=DEFAULT_CHUNK_SIZE,
                help="Choose a smaller chunk if the server runs out of memory.",
            )
            output_choice = st.radio(
                "Download format",
                ["CSV and XLSX", "CSV only", "XLSX only"],
                horizontal=True,
            )

            if st.button("Remove Duplicates", type="primary"):
                output_dir = tempfile.mkdtemp(prefix="csv_dedupe_")
                progress = st.progress(0, text="Starting...")
                status = st.empty()
                metrics = st.empty()
                final_stats = None
                try:
                    for stats in create_outputs(
                        uploaded_file, selected_column, output_dir, uploaded_file.name, chunk_size
                    ):
                        final_stats = stats
                        status.write(
                            f"Processed **{stats['processed']:,}** rows · "
                            f"Duplicates removed: **{stats['duplicates']:,}**"
                        )
                        metrics.write(
                            f"Rows kept: **{stats['kept']:,}** · Chunks completed: **{stats['chunk']:,}**"
                        )
                        progress.progress(min(0.95, stats["chunk"] / (stats["chunk"] + 2)))
                    progress.progress(1.0, text="Completed")
                    st.session_state["dedupe_output_dir"] = output_dir
                    st.session_state["dedupe_csv_path"] = final_stats["csv_path"]
                    st.session_state["dedupe_xlsx_path"] = final_stats["xlsx_path"]
                    st.session_state["dedupe_stats"] = final_stats
                    st.session_state["dedupe_output_choice"] = output_choice
                    st.success(
                        f"Completed. Removed {final_stats['duplicates']:,} duplicate rows; "
                        f"kept {final_stats['kept']:,} rows."
                    )
                except Exception as exc:
                    st.error(f"Could not process this file: {exc}")
                    import shutil
                    shutil.rmtree(output_dir, ignore_errors=True)
    except Exception as exc:
        st.error(f"Could not read the file's columns: {exc}")

if st.session_state.get("dedupe_stats"):
    st.subheader("Download Results")
    stats = st.session_state["dedupe_stats"]
    st.write(
        f"**Rows processed:** {stats['processed']:,} · **Rows kept:** {stats['kept']:,} · "
        f"**Duplicates removed:** {stats['duplicates']:,}"
    )
    choice = st.session_state.get("dedupe_output_choice", "CSV and XLSX")
    csv_path = st.session_state.get("dedupe_csv_path")
    xlsx_path = st.session_state.get("dedupe_xlsx_path")
    col1, col2 = st.columns(2)
    if choice in ("CSV and XLSX", "CSV only") and csv_path and os.path.isfile(csv_path):
        with col1:
            with open(csv_path, "rb") as handle:
                st.download_button(
                    "Download CSV", data=handle, file_name=Path(csv_path).name,
                    mime="text/csv", use_container_width=True,
                )
    if choice in ("CSV and XLSX", "XLSX only") and xlsx_path and os.path.isfile(xlsx_path):
        with col2:
            with open(xlsx_path, "rb") as handle:
                st.download_button(
                    "Download XLSX", data=handle, file_name=Path(xlsx_path).name,
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    use_container_width=True,
                )
